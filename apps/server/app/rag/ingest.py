"""File parsing + chunking.

Chunking strategy: split on markdown headings / paragraph boundaries first,
then pack greedily to ~chunk_size tokens with overlap. Heading context is
prepended to each chunk so retrieval hits stay self-describing — this matters
for voice: the agent reads chunks aloud-adjacent, so a chunk that starts
mid-thought produces a rambling answer.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from ..config import settings

# ~4 chars per token is close enough for packing decisions
CHARS_PER_TOKEN = 4


@dataclass
class Chunk:
    text: str
    heading: str
    index: int


def extract_text(filename: str, data: bytes, mime: str | None = None) -> str:
    name = filename.lower()
    if name.endswith(".pdf") or (mime or "").endswith("pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages)
    if name.endswith(".docx") or "wordprocessingml" in (mime or ""):
        import docx

        document = docx.Document(io.BytesIO(data))
        parts: list[str] = []
        for para in document.paragraphs:
            style = (para.style.name or "").lower() if para.style else ""
            text = para.text.strip()
            if not text:
                continue
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "2"
                parts.append("#" * int(level) + " " + text)
            else:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n\n".join(parts)
    # txt / md / anything text-like
    return data.decode("utf-8", errors="replace")


def _split_blocks(text: str) -> list[tuple[str, str]]:
    """Return [(heading_context, block_text)] split on headings/blank lines."""
    lines = text.splitlines()
    blocks: list[tuple[str, str]] = []
    heading = ""
    buf: list[str] = []

    def flush():
        nonlocal buf
        block = "\n".join(buf).strip()
        if block:
            blocks.append((heading, block))
        buf = []

    for line in lines:
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            flush()
            heading = m.group(2).strip()
            continue
        if not line.strip():
            if buf:
                buf.append("")
            continue
        buf.append(line)
    flush()

    # further split blocks into paragraphs
    out: list[tuple[str, str]] = []
    for h, block in blocks:
        for para in re.split(r"\n\s*\n", block):
            para = para.strip()
            if para:
                out.append((h, para))
    return out


def chunk_text(text: str) -> list[Chunk]:
    max_chars = settings.chunk_size_tokens * CHARS_PER_TOKEN
    overlap_chars = settings.chunk_overlap_tokens * CHARS_PER_TOKEN

    paras = _split_blocks(text)
    chunks: list[Chunk] = []
    cur: list[str] = []
    cur_heading = ""
    cur_len = 0

    def flush():
        nonlocal cur, cur_len
        if not cur:
            return
        body = "\n\n".join(cur).strip()
        text_out = f"{cur_heading}\n\n{body}" if cur_heading else body
        chunks.append(Chunk(text=text_out, heading=cur_heading, index=len(chunks)))
        # keep tail as overlap for continuity across chunk boundaries
        tail = body[-overlap_chars:] if overlap_chars and len(body) > overlap_chars else ""
        cur = [tail] if tail else []
        cur_len = len(tail)

    for heading, para in paras:
        if heading != cur_heading and cur:
            flush()
            cur = []  # heading change: don't carry overlap across sections
            cur_len = 0
        cur_heading = heading
        # hard-split single paragraphs that exceed the budget
        while len(para) > max_chars:
            cut = para.rfind(". ", 0, max_chars)
            cut = cut + 1 if cut > max_chars // 2 else max_chars
            cur.append(para[:cut].strip())
            flush()
            para = para[cut:].strip()
        if cur_len + len(para) > max_chars:
            flush()
        if para:
            cur.append(para)
            cur_len += len(para)
    flush()
    return [c for c in chunks if c.text.strip()]
